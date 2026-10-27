#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把 docs/code-map-and-ownership.md 转成适合打印/给老师看的 DOCX。"""
import argparse
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_header_repeat(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_runs(paragraph, text):
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**") and token.endswith("**"):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        elif token.startswith("`") and token.endswith("`"):
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x99, 0x33, 0x00)
        elif token.startswith("["):
            mm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            label, url = mm.group(1), mm.group(2)
            if url.startswith("http"):
                add_hyperlink(paragraph, label, url)
            else:
                paragraph.add_run(label)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(8.5)
        if i < len(lines) - 1:
            r.add_break()


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip()
        cells = [c.strip() for c in raw.strip("|").split("|")]
        rows.append(cells)
        i += 1
    # 第二行是 |---|---|，去掉
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows = [rows[0]] + rows[2:]
    return rows, i


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = table.cell(ri, ci)
            text = row[ci] if ci < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            # 表格内保留简单行内代码/链接，去掉过长的 http 链接
            text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
            add_runs(p, text)
            for r in p.runs:
                r.font.size = Pt(8)
            if ri == 0:
                for r in p.runs:
                    r.bold = True
                set_cell_shading(cell, "D9E2F3")
        if ri == 0:
            set_table_header_repeat(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_footer_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run.font.size = Pt(9)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color in [("Heading 1", 18, "1F3864"), ("Heading 2", 14, "2E74B5"), ("Heading 3", 11.5, "404040")]:
        st = doc.styles[name]
        st.font.name = "微软雅黑"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(4)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input_md")
    ap.add_argument("output_docx")
    args = ap.parse_args()

    with open(args.input_md, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    style_document(doc)
    add_footer_page_number(doc.sections[0])

    i = 0
    first_h1_handled = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, code_lines)
            continue

        if stripped.startswith("|"):
            rows, nxt = parse_table(lines, i)
            add_table(doc, rows)
            i = nxt
            continue

        if stripped == "---":
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            title = re.sub(r"\*\*", "", m.group(2)).strip()
            if level == 1 and not first_h1_handled:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(title)
                r.font.name = "微软雅黑"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                r.font.size = Pt(22)
                r.bold = True
                r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
                p.paragraph_format.space_after = Pt(6)
                first_h1_handled = True
            else:
                doc.add_heading(title, level=min(level, 3))
            i += 1
            continue

        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, text)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            i += 1
            continue

        m = re.match(r"^\s*[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.add_run(m.group(1) + ". ").bold = True
            add_runs(p, m.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.core_properties.title = "校园二手交易系统 — 源码位置与团队分工说明"
    doc.core_properties.subject = "JavaWeb 课程项目答辩材料"
    doc.core_properties.author = "校园二手交易系统项目组"
    doc.save(args.output_docx)
    print(f"saved: {args.output_docx}")


if __name__ == "__main__":
    main()
